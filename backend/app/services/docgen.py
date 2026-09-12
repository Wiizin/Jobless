"""Renders structured GeneratedDocument content to .docx via python-docx.

Only called after human approval per the workflow (routers/documents.py
sets status=approved before calling render_*). Nothing here submits or
sends the file anywhere — it just writes a file to disk for the user to
download.
"""
from __future__ import annotations

import os
import uuid

from docx import Document as DocxDocument
from docx.shared import Pt

from app.config import get_settings
from app.db.models import Document as DocumentModel
from app.db.models import Experience, Offer, Profile

settings = get_settings()


def _ensure_output_dir() -> str:
    os.makedirs(settings.generated_docs_dir, exist_ok=True)
    return settings.generated_docs_dir


def render_cv(profile: Profile, offer: Offer, document: DocumentModel, experiences: list[Experience]) -> str:
    """Render a CV emphasizing the selected experiences/bullets. Returns file path."""
    doc = DocxDocument()

    doc.add_heading(profile.full_name, level=0)
    contact_line = " | ".join(filter(None, [profile.email, profile.phone, profile.location]))
    doc.add_paragraph(contact_line)

    if profile.summary:
        doc.add_heading("Summary", level=1)
        doc.add_paragraph(profile.summary)

    doc.add_heading("Experience", level=1)
    exp_by_id = {e.id: e for e in experiences}
    for exp_id_str in document.selected_experiences:
        exp = exp_by_id.get(uuid.UUID(exp_id_str) if isinstance(exp_id_str, str) else exp_id_str)
        if exp is None:
            continue
        p = doc.add_paragraph()
        run = p.add_run(f"{exp.title} — {exp.company}")
        run.bold = True
        run.font.size = Pt(12)
        if exp.location or exp.start_date or exp.end_date:
            meta = f"{exp.location or ''} ({exp.start_date or ''} – {exp.end_date or 'Present' if exp.is_current else exp.end_date or ''})"
            doc.add_paragraph(meta)

    if document.bullets:
        doc.add_heading("Highlights", level=1)
        for bullet in document.bullets:
            doc.add_paragraph(bullet, style="List Bullet")

    if profile.education:
        doc.add_heading("Education", level=1)
        for edu in profile.education:
            doc.add_paragraph(f"{edu.degree or ''} in {edu.field_of_study or ''} — {edu.institution}")

    if profile.certifications:
        doc.add_heading("Certifications", level=1)
        for cert in profile.certifications:
            doc.add_paragraph(f"{cert.name} ({cert.issuer or ''})")

    out_dir = _ensure_output_dir()
    filename = f"cv_{offer.id}_{document.id}.docx"
    path = os.path.join(out_dir, filename)
    doc.save(path)
    return path


def render_cover_letter(profile: Profile, offer: Offer, document: DocumentModel) -> str:
    """Render the cover letter sections to .docx. Returns file path."""
    doc = DocxDocument()
    doc.add_heading(f"Cover Letter — {offer.title} at {offer.company}", level=1)
    doc.add_paragraph(profile.full_name)
    doc.add_paragraph("")

    sections = document.cover_letter_sections or {}
    intro = sections.get("intro")
    if intro:
        doc.add_paragraph(intro)
    for para in sections.get("body_paragraphs", []):
        doc.add_paragraph(para)
    closing = sections.get("closing")
    if closing:
        doc.add_paragraph(closing)

    out_dir = _ensure_output_dir()
    filename = f"cover_letter_{offer.id}_{document.id}.docx"
    path = os.path.join(out_dir, filename)
    doc.save(path)
    return path
